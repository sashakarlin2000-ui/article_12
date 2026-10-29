import csv
import os

import docx
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC = "article_v1.docx"
d = docx.Document(DOC)
paras = d.paragraphs


def find(prefix):
    for p in paras:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)


def set_text(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)


def insert_after(anchor, text="", style=None):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if style:
        para.style = style
    if text:
        para.add_run(text)
    return para


def read_eval(path):
    with open(path, encoding="utf-8-sig") as f:
        return next(csv.DictReader(f))


def f3(x):
    return f"{float(x):.3f}"


def f4(x):
    return f"{float(x):.4f}"


# ---------- status ----------
set_text(
    find("Статус рукописи"),
    "Статус рукописи: рабочая исследовательская версия. Приведены предварительные "
    "эмпирические результаты ограниченного прогона (равный набор наблюдений, один "
    "актив, три LLM). Полный протокол E1–E6 на шести активах, walk-forward по "
    "нескольким окнам и интервалы неопределённости остаются задачей следующих итераций.",
)

# ---------- RU abstract ----------
set_text(
    find("В работе предлагается экспериментальная схема"),
    "В работе исследуется схема прогнозирования направления движения фондового рынка, "
    "объединяющая определение рыночного режима, поиск исторических аналогов и "
    "агрегирование прогнозов нескольких больших языковых моделей (LLM). "
    "В предварительном вневыборочном прогоне на дневных данных AAPL (50 торговых дней, "
    "16 июля — 7 октября 2026 г.) сравнивались три модели Google (Gemini 3.1 Flash Lite, "
    "Gemini 3.5 Flash Lite, Gemma 4 31B) при обычном и режимно-зависимом retrieval (top-3) "
    "и два правила ансамбля. Лучшую Accuracy показала Gemma 4 31B с обычным retrieval "
    "(0.480). Ансамбль большинством дал Accuracy 0.420, лучший MCC (0.188) и накопленную "
    "доходность 12.5% (Sharpe 2.60). Режимно-зависимый retrieval ухудшил качество: "
    "Accuracy ансамбля 0.306, доходность 0.5%. Полученные данные не поддерживают H2 и "
    "дают лишь частичную поддержку H3; H1 и H4 в этом прогоне не проверялись. Результаты "
    "ограничены одним активом и коротким окном и рассматриваются как предварительные.",
)

# ---------- EN abstract ----------
set_text(
    find("This paper proposes an experimental framework"),
    "This paper studies a forecasting framework for stock-market direction that combines "
    "market-regime detection, retrieval of historical analogues, and aggregation of forecasts "
    "from multiple large language models (LLMs). A preliminary out-of-sample run on daily AAPL "
    "data (50 trading days, 16 July – 7 October 2026) compared three Google models "
    "(Gemini 3.1 Flash Lite, Gemini 3.5 Flash Lite, Gemma 4 31B) under standard and "
    "regime-aware top-3 retrieval and two ensemble rules. Gemma 4 31B with standard retrieval "
    "achieved the highest accuracy (0.480). The majority ensemble reached accuracy 0.420, the "
    "best MCC (0.188), and a 12.5% cumulative return (Sharpe 2.60). Regime-aware retrieval "
    "degraded performance: ensemble accuracy fell to 0.306 and return to 0.5%. The results do "
    "not support H2 and provide only partial support for H3; H1 and H4 were not tested in this "
    "run. Findings are limited to one asset and a short window and should be treated as "
    "preliminary.",
)

# ---------- methodology -> actual run ----------
set_text(
    find("До запуска финального тестирования"),
    "Эксперимент использует дневные данные по шести активам: SPX, AAPL, MSFT, NVDA, AMZN и TSLA. "
    "Период наблюдений — с 26 июля 2017 г. по 9 октября 2026 г., по 2315 торговых дней на актив. "
    "Данные разделены хронологически: обучение — до 31 декабря 2022 г., валидация — с 3 января "
    "2023 г. по 30 июня 2025 г., финальный тест — с 1 июля 2025 г. Из-за лимитов бесплатных API "
    "финальный прогон ограничен последними 60 торговыми днями каждого актива, а для сопоставимости "
    "моделей анализируется равный набор наблюдений: AAPL, 50 дней для обычного retrieval и 49 дней "
    "для режимно-зависимого (16 июля — 7 октября 2026 г.). Выбор окна не зависел от результатов моделей.",
)

set_text(
    find("Для наблюдения на момент t"),
    "Целевая переменная для горизонта h=1 определяется по ценам закрытия: y(t)=1, если "
    "Close(t+1)>Close(t), и y(t)=0 в противном случае. Прогноз формируется по данным, доступным "
    "на закрытии дня t; сделка открывается по закрытию t и закрывается по закрытию t+1. "
    "Класс HOLD трактуется как отсутствие позиции.",
)

set_text(
    find("Потенциальные признаки включают"),
    "В прогоне использованы признаки: доходность за 5 дней, волатильность за 20 дней, RSI(14) "
    "и расстояние цены до MA200; для расчёта сходства дополнительно использовались волатильность "
    "за 5 дней и расстояние до MA20. Все преобразования, включая стандартизацию, оценивались "
    "только на доступной к моменту t истории.",
)

set_text(
    find("В качестве кандидатов рассматриваются K-Means"),
    "Рыночные режимы определены методом K-Means (k=4) по признакам return_5d, volatility_20d, "
    "distance_from_MA20 и distance_from_MA200. Модель обучена на данных до 31 декабря 2022 г., "
    "метки сглажены окном 3 дня. Кластерам присвоены интерпретируемые названия Bull, Bear, "
    "Sideways и High Volatility по их статистикам. В анализируемом окне AAPL преобладает "
    "Sideways (39 из 50 дней), реже встречаются Bull (7) и Bear (4); High Volatility не встретился.",
)

set_text(
    find("Для каждого прогнозного момента формируется вектор"),
    "Для каждого прогнозного момента формировался вектор признаков и выполнялся поиск top-3 "
    "наиболее похожих исторических окон по косинусному расстоянию. Признаки нормализовались "
    "параметрами, оценёнными только на доступной истории; все аналоги ограничены датами, "
    "предшествующими t, а окно их исхода заканчивалось до t.",
)

set_text(
    find("Каждая выбранная LLM получает стандартизированный вход"),
    "Каждая LLM получала стандартизированный вход: текущие признаки, метку режима, top-3 "
    "исторических аналога без их будущих исходов и требование вернуть JSON с полями prediction, "
    "confidence, probability_up и explanation. В прогоне использованы три модели Google: "
    "Gemini 3.1 Flash Lite, Gemini 3.5 Flash Lite и Gemma 4 31B (temperature 0.2). Ответы "
    "кэшировались, сырые ответы сохранялись; невалидный JSON не заменялся случайным сигналом.",
)

set_text(
    find("Сравниваются следующие стратегии объединения"),
    "Сравнивались два правила объединения: простое большинство голосов и взвешивание по "
    "confidence. Режимно-зависимые веса и ансамбль с обученными весами в текущем прогоне "
    "не запускались, поэтому E6 остаётся открытым.",
)

set_text(
    find("Для воспроизводимости фиксируются"),
    "Для воспроизводимости зафиксированы названия моделей, дата запуска, промпт, temperature=0.2, "
    "формат ответа и правила обработки ошибок. Полные сырые ответы и журнал вызовов сохранены в "
    "llm_responses.jsonl и results/llm/llm_prompts_log.jsonl. Бесплатные модели OpenRouter "
    "(Nemotron 3.5, Liquid LFM 2.5, Gemma 4 31B) были исключены из равного анализа, поскольку "
    "дневной лимит free-тира (50 запросов) был исчерпан.",
)

# ---------- section 8 ----------
h8 = find("8. Результаты")
old1 = find("Раздел предназначен для заполнения")
old2 = find("До получения эмпирических результатов")

last = h8
last = insert_after(last, "8.1. Данные и вычислительный протокол", "Heading 2")
last = insert_after(
    last,
    "Эмпирическая часть выполнена на дневных данных AAPL за 16 июля — 7 октября 2026 г. "
    "(50 торговых дней; для режимно-зависимого retrieval доступно 49 дней). Доля дней роста "
    "составила 52%. Тест является вневыборочным: все аналоги и параметры режимной модели "
    "предшествуют прогнозному моменту. Классификация оценивалась по Accuracy, Balanced Accuracy, "
    "F1, MCC и Brier score; торговая симуляция — с комиссией 0.1%, проскальзыванием 0.05%, "
    "длинными и короткими позициями, HOLD означает отсутствие позиции. Для сравнения: "
    "Buy & Hold даёт 1.65% за окно standard и -3.03% за окно regime-aware.",
)
last = insert_after(last, "8.2. Распределение режимов", "Heading 2")
last = insert_after(
    last,
    "В окне преобладает Sideways — 39 из 50 дней; Bull — 7 дней, Bear — 4 дня; режим High "
    "Volatility не встретился. Поэтому выводы о поведении в разных режимах носят предварительный "
    "характер: наблюдений в Bull и Bear недостаточно для устойчивых оценок.",
)
last = insert_after(last, "8.3. Классификационные результаты", "Heading 2")
last = insert_after(
    last,
    "Таблица 1 содержит метрики для отдельных моделей и ансамблей. Лучшую Accuracy показала "
    "Gemma 4 31B с обычным retrieval — 0.480 (Balanced Accuracy 0.474, MCC 0.170). Ансамбль "
    "большинством с обычным retrieval дал Accuracy 0.420 и лучший среди конфигураций MCC 0.188. "
    "Gemini 3.1 Flash Lite показал более низкую Accuracy (0.320), но сопоставимый MCC (0.188) "
    "и высокую доходность за счёт меньшего числа сделок. Gemini 3.5 Flash Lite оказался "
    "слабейшим: 0.280 при обычном retrieval и 0.163 при режимно-зависимом.",
)
cap1 = insert_after(last, "Таблица 1. Классификационные метрики, AAPL, равный набор наблюдений.")
for r in cap1.runs:
    r.bold = True
last = insert_after(cap1, "")
last = insert_after(last, "8.4. Торговые результаты", "Heading 2")
last = insert_after(
    last,
    "Таблица 2 показывает результаты торговой симуляции. Наибольшую накопленную доходность при "
    "обычном retrieval дали Gemini 3.1 Flash Lite (16.6%, Sharpe 6.18) и Gemma 4 31B (12.9%, "
    "Sharpe 2.47); ансамбль большинством — 12.5% (Sharpe 2.60). Режимно-зависимый retrieval "
    "ухудшил результаты: ансамбль majority дал 0.5% (Sharpe 0.22), confidence — -0.05% "
    "(Sharpe 0.10). Высокий Sharpe у Gemini 3.1 объясняется малым числом сделок (15) и низкой "
    "волатильностью, а не высокой точностью, поэтому не должен интерпретироваться как "
    "устойчивое преимущество.",
)
cap2 = insert_after(last, "Таблица 2. Торговые метрики, AAPL, издержки учтены.")
for r in cap2.runs:
    r.bold = True
last = insert_after(cap2, "")
last = insert_after(last, "8.5. Проверка гипотез", "Heading 2")
hyp = [
    "H1 (retrieval против отсутствия retrieval) не проверялась: конфигурация E2 без retrieval "
    "в текущем прогоне не запускалась.",
    "H2 (режимно-зависимый retrieval лучше обычного) не получила поддержки: у всех трёх моделей "
    "и у обоих ансамблей метрики при regime-aware ниже. Accuracy ансамбля снизилась с 0.420 до "
    "0.306, MCC — с 0.188 до 0.045, накопленная доходность — с 12.5% до 0.5%.",
    "H3 (ансамбль лучше отдельных моделей) поддержана частично: ансамбль дал лучший MCC (0.188) "
    "и конкурентный Sharpe (2.60), но уступил Gemma 4 31B по Accuracy (0.420 против 0.480). "
    "Устойчивость по временным окнам не оценивалась.",
    "H4 (режимно-зависимые веса) не проверялась: E6 в текущем прогоне отсутствует.",
    "Согласие всех трёх моделей по направлению наблюдалось лишь в 34% дней, что объясняет "
    "расхождение между majority и confidence-weighted ансамблями.",
]
for h in hyp:
    last = insert_after(last, h, "List Bullet")
last = insert_after(last, "8.6. Ограничения текущего прогона", "Heading 2")
last = insert_after(
    last,
    "Результаты ограничены одним активом (AAPL), 50 торговыми днями и тремя моделями Google. "
    "Дневные лимиты бесплатных API не позволили включить OpenRouter-модели (Nemotron, Liquid, "
    "Gemma через OpenRouter) и полный набор шести активов. Интервалы неопределённости (блочный "
    "бутстрэп), walk-forward по нескольким окнам, абляции E1/E2/E6 и коррекция на множественные "
    "сравнения не выполнялись. Поэтому приведённые числа следует рассматривать как предварительные "
    "и не использовать для утверждений о превосходстве метода.",
)

# remove placeholders
for p in (old1, old2):
    p._element.getparent().remove(p._element)

# ---------- tables ----------
configs = [
    ("eval_single_gemma-4-31b-it_standard", "Gemma 4 31B — standard"),
    ("eval_single_gemini-3.1-flash-lite_standard", "Gemini 3.1 Flash Lite — standard"),
    ("eval_single_gemini-3.5-flash-lite_standard", "Gemini 3.5 Flash Lite — standard"),
    ("eval_single_gemma-4-31b-it_regime-aware", "Gemma 4 31B — regime-aware"),
    ("eval_single_gemini-3.1-flash-lite_regime-aware", "Gemini 3.1 Flash Lite — regime-aware"),
    ("eval_single_gemini-3.5-flash-lite_regime-aware", "Gemini 3.5 Flash Lite — regime-aware"),
    ("eval_standard_majority_voting", "Ансамбль majority — standard"),
    ("eval_standard_confidence_weighted_voting", "Ансамбль confidence — standard"),
    ("eval_regime-aware_majority_voting", "Ансамбль majority — regime-aware"),
    ("eval_regime-aware_confidence_weighted_voting", "Ансамбль confidence — regime-aware"),
]

base = "results_equal"
cls_header = ["Конфигурация", "n", "Accuracy", "Bal. Acc.", "F1", "MCC", "Brier"]
cls_rows = [cls_header]
trd_header = ["Конфигурация", "Доходность", "Sharpe", "Max DD", "Win rate", "Сделки"]
trd_rows = [trd_header]
for folder, label in configs:
    c = read_eval(os.path.join(base, folder, "results_classification.csv"))
    t = read_eval(os.path.join(base, folder, "results_trading.csv"))
    cls_rows.append([
        label, c["n_days"], f3(c["accuracy"]), f3(c["balanced_accuracy"]),
        f3(c["f1"]), f3(c["mcc"]), f3(c["brier_score"]),
    ])
    trd_rows.append([
        label, f"{float(t['cumulative_return']) * 100:.2f}%", f"{float(t['sharpe_ratio']):.2f}",
        f"{float(t['max_drawdown']) * 100:.2f}%", f3(t["win_rate"]), t["n_trades"],
    ])


def add_table_after(anchor, rows):
    table = d.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            if i == 0:
                run.bold = True
    anchor._p.addnext(table._tbl)
    return table


add_table_after(cap1, cls_rows)
add_table_after(cap2, trd_rows)

# ---------- discussion ----------
set_text(
    find("Интерпретация результатов должна учитывать"),
    "Интерпретация результатов должна учитывать, что рыночные режимы могут быть нестабильными, "
    "а их границы зависят от признаков и алгоритма. Полученные предварительные результаты "
    "согласуются с предупреждением Kirby [2]: наличие режимной структуры не гарантирует выигрыша "
    "от её использования. В нашем окне фильтрация аналогов по режиму сократила пул кандидатов "
    "и ухудшила качество, особенно для Gemini 3.5 Flash Lite. Это может быть следствием как "
    "малого числа наблюдений в Bull и Bear, так и нестабильности границ режимов. LLM чувствительны "
    "к формулировкам промптов и нестабильны между запусками; в текущем прогоне согласие всех трёх "
    "моделей наблюдалось лишь в 34% дней. Дополнительные ограничения связаны с выбором актива и "
    "периода, качеством данных, транзакционными издержками, возможной переоптимизацией весов и "
    "множественными экспериментальными сравнениями.",
)

# ---------- conclusion ----------
set_text(
    find("В работе сформулирован воспроизводимый план"),
    "В работе сформулирован и частично выполнен воспроизводимый план исследования режимно-зависимого "
    "retrieval и ансамблирования LLM. В предварительном вневыборочном прогоне на AAPL (50 дней, три "
    "модели Google) обычный retrieval оказался лучше режимно-зависимого: ансамбль большинством дал "
    "Accuracy 0.420, MCC 0.188 и накопленную доходность 12.5% против 0.306, 0.045 и 0.5% при "
    "regime-aware. Лучшая одиночная модель по Accuracy — Gemma 4 31B (0.480); ансамбль не превзошёл "
    "её по Accuracy, но дал лучший MCC и конкурентный Sharpe. Гипотеза H2 не поддержана, H3 "
    "поддержана частично; H1 и H4 остаются непроверенными. Полные выводы требуют расширения на шесть "
    "активов, конфигураций E1/E2/E6, нескольких walk-forward окон и интервалов неопределённости.",
)

d.save(DOC)
print("saved", DOC)
